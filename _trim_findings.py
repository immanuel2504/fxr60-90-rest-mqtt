from copy import copy
from pathlib import Path

from openpyxl import load_workbook

DST = Path(__file__).resolve().parent / "FXR-API-Findings-For-Developer_2026-09-07.xlsx"


def main():
    wb = load_workbook(DST)
    ws = wb["Questions"]
    keep = []
    for row in range(3, ws.max_row + 1):
        status = (ws.cell(row, 9).value or "").strip()
        if status == "Updated":
            continue
        keep.append(row)

    snapshot = []
    for row in keep:
        cells = []
        for col in range(1, 11):
            c = ws.cell(row, col)
            cells.append(
                {
                    "value": c.value,
                    "font": copy(c.font),
                    "fill": copy(c.fill),
                    "alignment": copy(c.alignment),
                    "border": copy(c.border),
                    "number_format": c.number_format,
                    "height": ws.row_dimensions[row].height,
                }
            )
        snapshot.append(cells)

    for row in range(3, ws.max_row + 1):
        for col in range(1, 11):
            ws.cell(row, col).value = None

    for i, cells in enumerate(snapshot):
        dest_row = 3 + i
        cells[0]["value"] = i + 1
        if cells[8]["height"]:
            ws.row_dimensions[dest_row].height = cells[8]["height"]
        for col, data in enumerate(cells, start=1):
            c = ws.cell(dest_row, col)
            c.value = data["value"]
            c.font = data["font"]
            c.fill = data["fill"]
            c.alignment = data["alignment"]
            c.border = data["border"]
            c.number_format = data["number_format"]

    ws.cell(1, 1).value = (
        "Column I is the schema check vs openAPISpec 11.yaml (schema only). "
        "Dated 2026-09-07. Updated (fixed) rows were removed. "
        "Column H is empty on new rows — for the developer reply."
    )
    wb.save(DST)
    print("kept", len(snapshot), "open rows")
    for cells in snapshot:
        print(cells[0]["value"], cells[3]["value"], "|", cells[8]["value"])


if __name__ == "__main__":
    main()
