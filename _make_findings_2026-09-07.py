from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "FXR-API-Findings-For-Developer_2026-09-feedback .xlsx"
DST = ROOT / "FXR-API-Findings-For-Developer_2026-09-07.xlsx"


def main():
    wb = load_workbook(SRC)
    ws = wb["Questions"]
    r = 28
    values = [
        26,
        "PUT /cloud/impinjGen2X",
        "set_impinjGen2X",
        "only one feature per request",
        (
            "Any two of fastID, tagProtect, tagFocus, tagQuieting in one request "
            "return HTTP 422: Only one Impinj Gen2X feature can be configured at a "
            "time. All four are exclusive. Tested on reader 5.0.7 (five combinations)."
        ),
        (
            "v3.0.0 / spec 11 removed the mutually exclusive table. The PUT "
            "description does not say only one feature is allowed. An older spec "
            "said only fastID, tagFocus, tagQuieting conflict and tagProtect could "
            "be mixed. That is not what the reader does."
        ),
        (
            "Can you document that exactly one of fastID, tagProtect, tagFocus, or "
            "tagQuieting is allowed per request? Combining any two is rejected. "
            "Please restore a mutual-exclusion note covering all four features."
        ),
        None,
        "Not in spec 11",
        (
            "Live 2026-09-06. Spec 11 PUT description mentions applyImpinjGen2X "
            "but no one-feature rule. Schema allows more than one property "
            "(minProperties: 1 only)."
        ),
    ]
    for col, value in enumerate(values, start=1):
        src_c = ws.cell(27, col)
        dst_c = ws.cell(r, col)
        dst_c.value = value
        if src_c.has_style:
            dst_c.font = copy(src_c.font)
            dst_c.alignment = copy(src_c.alignment)
            dst_c.border = copy(src_c.border)
            dst_c.number_format = src_c.number_format
            dst_c.fill = copy(src_c.fill)
    ws.cell(r, 9).fill = PatternFill(
        start_color="FFFF00", end_color="FFFF00", fill_type="solid"
    )
    if ws.row_dimensions[27].height:
        ws.row_dimensions[r].height = ws.row_dimensions[27].height

    ws.cell(1, 1).value = (
        "Column I is the schema check vs openAPISpec 11.yaml (schema only). "
        "Dated 2026-09-07. Filter I for Not in spec 11 (new live findings) or "
        "Said updated — not in schema. Updated rows match the schema. "
        "Column H is empty on new rows — for the developer reply."
    )
    wb.save(DST)
    print("wrote", DST)


if __name__ == "__main__":
    main()
